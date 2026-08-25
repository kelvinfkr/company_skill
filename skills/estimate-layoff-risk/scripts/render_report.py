#!/usr/bin/env python3
"""Typeset a layoff-risk economics report from structured JSON.

Pipeline: report.json -> DOCX (python-docx) -> PDF (LibreOffice).
Markdown is not used for final report rendering.
"""
from __future__ import annotations
import argparse, json, os, shutil, subprocess, tempfile
from pathlib import Path
from typing import Any, Iterable

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt, RGBColor

FONT_CN = "Noto Sans CJK SC"
FONT_SERIF = FONT_CN
INK = "18212B"
MUTED = "667085"
LIGHT = "F2F4F7"
LINE = "D0D5DD"
DARK = "101828"
ACCENT = "344054"
SOFT = "F8F9FB"


def set_cell_shading(cell, fill: str):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        tcPr.append(shd)
    shd.set(qn('w:fill'), fill)


def set_cell_border(cell, color=LINE, size=4):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    borders = tcPr.first_child_found_in('w:tcBorders')
    if borders is None:
        borders = OxmlElement('w:tcBorders')
        tcPr.append(borders)
    for edge in ('top','left','bottom','right','insideH','insideV'):
        tag = 'w:' + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn('w:val'), 'single')
        element.set(qn('w:sz'), str(size))
        element.set(qn('w:color'), color)


def set_repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr()
    tblHeader = OxmlElement('w:tblHeader')
    tblHeader.set(qn('w:val'), 'true')
    trPr.append(tblHeader)


def prevent_row_split(row):
    trPr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement('w:cantSplit')
    trPr.append(cant_split)


def set_cell_margins(cell, top=80, start=90, bottom=80, end=90):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in('w:tcMar')
    if tcMar is None:
        tcMar = OxmlElement('w:tcMar')
        tcPr.append(tcMar)
    for m, v in [('top',top),('start',start),('bottom',bottom),('end',end)]:
        node = tcMar.find(qn('w:'+m))
        if node is None:
            node = OxmlElement('w:'+m)
            tcMar.append(node)
        node.set(qn('w:w'), str(v)); node.set(qn('w:type'),'dxa')


def set_repeat_header_footer_font(paragraph, size=8, color=MUTED):
    for run in paragraph.runs:
        run.font.name = FONT_CN
        run._element.rPr.rFonts.set(qn('w:eastAsia'), FONT_CN)
        run.font.size = Pt(size)
        run.font.color.rgb = RGBColor.from_string(color)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("第 ")
    fldChar1 = OxmlElement('w:fldChar'); fldChar1.set(qn('w:fldCharType'), 'begin')
    instrText = OxmlElement('w:instrText'); instrText.set(qn('xml:space'), 'preserve'); instrText.text = 'PAGE'
    fldChar2 = OxmlElement('w:fldChar'); fldChar2.set(qn('w:fldCharType'), 'end')
    run._r.append(fldChar1); run._r.append(instrText); run._r.append(fldChar2)
    paragraph.add_run(" 页")
    set_repeat_header_footer_font(paragraph)


def configure_doc(doc: Document):
    sec = doc.sections[0]
    sec.page_width = Mm(210); sec.page_height = Mm(297)
    sec.different_first_page_header_footer = True
    sec.top_margin = Mm(14); sec.bottom_margin = Mm(13)
    sec.left_margin = Mm(18); sec.right_margin = Mm(18)
    sec.header_distance = Mm(8); sec.footer_distance = Mm(8)

    styles = doc.styles
    normal = styles['Normal']
    normal.font.name = FONT_CN; normal._element.rPr.rFonts.set(qn('w:eastAsia'), FONT_CN)
    normal.font.size = Pt(9.5); normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(4.5); normal.paragraph_format.line_spacing = 1.22

    for nm, size, before, after, color in [
        ('Title', 25, 0, 8, DARK), ('Subtitle', 10.5, 0, 8, MUTED),
        ('Heading 1', 16, 16, 8, DARK), ('Heading 2', 12, 10, 5, ACCENT),
        ('Heading 3', 10.5, 8, 4, ACCENT)]:
        st=styles[nm]; st.font.name=FONT_CN; st._element.rPr.rFonts.set(qn('w:eastAsia'),FONT_CN)
        st.font.size=Pt(size); st.font.color.rgb=RGBColor.from_string(color)
        st.paragraph_format.space_before=Pt(before); st.paragraph_format.space_after=Pt(after)
        st.paragraph_format.keep_with_next=True


def text_run(p, text, *, bold=False, size=None, color=None, font=FONT_CN):
    r=p.add_run(str(text))
    r.bold=bold; r.font.name=font; r._element.rPr.rFonts.set(qn('w:eastAsia'),font)
    if size: r.font.size=Pt(size)
    if color: r.font.color.rgb=RGBColor.from_string(color)
    return r


def add_cover(doc: Document, data: dict):
    meta=data.get('meta',{})
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(78)
    text_run(p, meta.get('company','Company'), bold=True, size=27, color=DARK, font=FONT_SERIF)
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    text_run(p, '岗位非主动离开风险研究报告', bold=True, size=19, color=DARK)
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    text_run(p, meta.get('role',''), size=10.5, color=MUTED)

    doc.add_paragraph('')
    rows=[
        ('案例代号',meta.get('case_alias','-')),
        ('法律雇主',meta.get('legal_employer','-')),
        ('国家/地区',meta.get('country','-')),
        ('参考日期',meta.get('reference_date','-')),
        ('总体置信度',meta.get('confidence','-')),
        ('结果定义',meta.get('outcome_definition','-')),
    ]
    t=doc.add_table(rows=0, cols=2); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
    for k,v in rows:
        c=t.add_row().cells; c[0].width=Cm(3.1); c[1].width=Cm(9.7)
        set_cell_shading(c[0],LIGHT)
        for idx,txt in enumerate((k,v)):
            c[idx].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER; set_cell_margins(c[idx],120,140,120,140)
            p=c[idx].paragraphs[0]; p.paragraph_format.space_after=Pt(0)
            text_run(p,txt,bold=(idx==0),size=9.2,color=INK if idx else ACCENT)
    doc.add_paragraph('')
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    text_run(p,'方法说明：报告区分公开事实、候选人自述、派生计算、模型假设与推断。概率为情景区间，不是保证或法律意见。',size=8.2,color=MUTED)
    doc.add_page_break()


def add_header_footer(doc: Document, company: str):
    sec=doc.sections[0]
    hdr=sec.header.paragraphs[0]; hdr.alignment=WD_ALIGN_PARAGRAPH.LEFT
    text_run(hdr, f'{company}  |  Layoff Risk Economics', size=7.8, color=MUTED)
    ftr=sec.footer.paragraphs[0]; add_page_number(ftr)


def add_section_title(doc, num, title, note=None):
    p=doc.add_paragraph(style='Heading 1')
    text_run(p, f'{num:02d}', bold=True, size=9.5, color=MUTED)
    text_run(p, f'  {title}', bold=True, size=16, color=DARK)
    if note:
        q=doc.add_paragraph(); q.paragraph_format.space_after=Pt(7)
        text_run(q,note,size=8.8,color=MUTED)


def add_callout(doc, text):
    t=doc.add_table(rows=1,cols=1); t.autofit=False; t.alignment=WD_TABLE_ALIGNMENT.CENTER
    cell=t.cell(0,0); set_cell_shading(cell,SOFT); set_cell_margins(cell,160,180,160,180)
    p=cell.paragraphs[0]; p.paragraph_format.space_after=Pt(0)
    text_run(p,text,bold=True,size=10,color=DARK)
    doc.add_paragraph('').paragraph_format.space_after=Pt(2)


def add_kpi_cards(doc, metrics: list[dict]):
    if not metrics: return
    chunk=3
    for start in range(0,len(metrics),chunk):
        batch=metrics[start:start+chunk]
        t=doc.add_table(rows=1,cols=len(batch)); t.autofit=False; t.alignment=WD_TABLE_ALIGNMENT.CENTER
        widths=[Cm(5.2)]*len(batch)
        for i,m in enumerate(batch):
            c=t.cell(0,i); c.width=widths[i]; set_cell_shading(c,SOFT); set_cell_margins(c,160,150,140,150)
            p=c.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(3)
            text_run(p,m.get('value','-'),bold=True,size=15,color=DARK)
            p=c.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(0)
            text_run(p,m.get('label',''),size=8.2,color=MUTED)
        doc.add_paragraph('').paragraph_format.space_after=Pt(1)


def add_table(doc, headers: list[str], rows: Iterable[Iterable[Any]], widths: list[float]|None=None, font_size=8.0):
    rows=list(rows)
    if not headers: return
    t=doc.add_table(rows=1, cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
    if widths is None:
        widths=[16.0/len(headers)]*len(headers)
    hdr=t.rows[0]; set_repeat_table_header(hdr); prevent_row_split(hdr)
    for i,h in enumerate(headers):
        c=hdr.cells[i]; c.width=Cm(widths[i]); set_cell_shading(c,ACCENT); set_cell_margins(c,75,100,75,100)
        p=c.paragraphs[0]; p.paragraph_format.space_after=Pt(0)
        text_run(p,h,bold=True,size=font_size,color='FFFFFF')
    for ri,row in enumerate(rows):
        added_row=t.add_row(); prevent_row_split(added_row); cells=added_row.cells
        for i,val in enumerate(row):
            c=cells[i]; c.width=Cm(widths[i]); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(c,65,100,65,100); set_cell_border(c)
            if ri%2==1: set_cell_shading(c,'FAFAFA')
            p=c.paragraphs[0]; p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1.05
            text_run(p,'' if val is None else val,size=font_size,color=INK)
    doc.add_paragraph('').paragraph_format.space_after=Pt(1)


def rows_by_keys(items, keys):
    return [[x.get(k,'') for k in keys] for x in items]


def add_bullets(doc, items):
    for item in items or []:
        p=doc.add_paragraph(style='Normal'); p.paragraph_format.left_indent=Cm(0.45); p.paragraph_format.first_line_indent=Cm(-0.28)
        text_run(p,'• ',bold=True,color=ACCENT); text_run(p,item,size=9.2)


def build_docx(data: dict, out_docx: Path):
    doc=Document(); configure_doc(doc)
    company=data.get('meta',{}).get('company','Company')
    add_cover(doc,data); add_header_footer(doc,company)

    add_section_title(doc,1,'执行摘要','按期限和退出机制分开报告，避免把不续签冒充法定裁员。')
    ex=data.get('executive',{})
    if ex.get('diagnosis'): add_callout(doc,ex['diagnosis'])
    if ex.get('dominant_uncertainty'): add_bullets(doc,[f"主要不确定性：{ex['dominant_uncertainty']}"])
    add_table(doc,['期限','结构性裁撤','合同/不续','绩效退出','任一非主动离开'],rows_by_keys(ex.get('risk_by_horizon',[]),['horizon','structural_layoff','contract_or_nonrenewal','performance_exit','any_involuntary_exit']),[2.2,3.1,3.1,3.0,4.6],7.4)

    add_section_title(doc,2,'公司压力与行动机制','先判断生存、利润率、资本重配、整合还是需求容量问题。')
    pressure=data.get('company_pressure',{})
    if pressure.get('regime'): add_callout(doc,f"行动机制：{pressure['regime']}")
    if pressure.get('summary'): add_bullets(doc,[pressure['summary']])
    add_table(doc,['指标','范围/数值','类型','证据'],rows_by_keys(pressure.get('metrics',[]),['metric','value','type','evidence']),[3.6,3.1,2.8,6.5],7.7)

    add_section_title(doc,3,'利润池与资本配置','收入不是价值；优先看贡献利润、现金转换和战略期权。')
    add_table(doc,['利润池','经济状态','现金/贡献','期权/战略','管理层意图','风险'],rows_by_keys(data.get('profit_pools',[]),['name','economics','contribution','option_value','management_intent','risk']),[3.0,3.1,2.5,2.5,3.1,1.8],7.1)

    add_section_title(doc,4,'利润池到岗位的依赖链','评估裁掉岗位节省的可避免成本与损失的利润、能力和期权。')
    add_table(doc,['利润池','关键流程','团队/岗位','可替代性','价值受损时间','净风险'],rows_by_keys(data.get('dependency_graph',[]),['profit_pool','process','role','substitutability','time_to_damage','risk']),[2.8,3.1,3.2,2.5,2.5,1.9],7.0)

    add_section_title(doc,5,'领导与资源配置机制','只根据可观察的历史决策与治理激励建模，不做人格式判断。')
    lead=data.get('leadership_regime',{})
    if lead.get('summary'): add_callout(doc,lead['summary'])
    add_bullets(doc,lead.get('evidence',[]))

    add_section_title(doc,6,'候选人可观察信号','候选人自述与公开证据分开保存；传闻不得升级为事实。')
    signals=data.get('candidate_signals',{})
    if signals.get('risk_up'):
        doc.add_paragraph('风险上行',style='Heading 2'); add_bullets(doc,signals['risk_up'])
    if signals.get('protective'):
        doc.add_paragraph('保护因素',style='Heading 2'); add_bullets(doc,signals['protective'])
    if signals.get('ambiguous'):
        doc.add_paragraph('待确认',style='Heading 2'); add_bullets(doc,signals['ambiguous'])

    doc.add_page_break()
    add_section_title(doc,7,'情景分析','区间宽度来自关键假设，而不是用单一点隐藏不确定性。')
    add_table(doc,['情景','关键假设','6个月风险','含义'],rows_by_keys(data.get('scenarios',[]),['name','assumptions','risk_6m','interpretation']),[2.2,6.7,2.8,4.3],7.4)

    add_section_title(doc,8,'证据与反证','每个重大风险判断都主动搜索相反证据。')
    doc.add_paragraph('支持风险上升的证据',style='Heading 2'); add_bullets(doc,[x if isinstance(x,str) else x.get('claim','') for x in data.get('evidence_for',[])])
    doc.add_paragraph('保护性证据或反证',style='Heading 2'); add_bullets(doc,[x if isinstance(x,str) else x.get('claim','') for x in data.get('evidence_against',[])])

    add_section_title(doc,9,'后续观察清单','说明哪些新事件会实质改变估计。')
    watch=data.get('watchlist',{})
    doc.add_paragraph('风险上行触发器',style='Heading 2'); add_bullets(doc,watch.get('risk_up',[]))
    doc.add_paragraph('风险下行触发器',style='Heading 2'); add_bullets(doc,watch.get('risk_down',[]))

    add_section_title(doc,10,'局限、隐私与来源','本报告服务候选人的个人规划，不用于替雇主选择裁撤对象。')
    add_bullets(doc,data.get('limitations',[]))
    if data.get('sources'):
        doc.add_paragraph('主要来源',style='Heading 2')
        add_bullets(doc,[x if isinstance(x,str) else f"{x.get('title','')}：{x.get('url','')}" for x in data['sources']])
    add_bullets(doc,['隐私说明：仅使用公开资料和候选人自愿提供的信息；候选人自述按授权范围脱敏。','免责声明：本报告为情景化职业风险分析，不构成就业、劳动法、投资或其他专业意见。'])

    doc.save(out_docx)


def convert_to_pdf(docx_path: Path, pdf_path: Path):
    with tempfile.TemporaryDirectory(prefix='cte_lo_') as td:
        env=os.environ.copy(); env['HOME']=td
        font_dir = Path(__file__).resolve().parent.parent / 'assets' / 'fonts'
        if font_dir.exists():
            private_fonts = Path(td) / '.fonts'
            shutil.copytree(font_dir, private_fonts)
            env['SAL_PRIVATE_FONTPATH'] = str(private_fonts)
            fc_cache = shutil.which('fc-cache')
            if fc_cache:
                subprocess.run(
                    [fc_cache, '-f', str(private_fonts)], check=True,
                    stdout=subprocess.PIPE, stderr=subprocess.PIPE, env=env,
                )
        outdir=Path(td)/'out'; outdir.mkdir()
        office = shutil.which('libreoffice') or shutil.which('soffice')
        if office is None:
            raise RuntimeError('LibreOffice/soffice executable was not found')
        cmd=[office,'--headless','--convert-to','pdf','--outdir',str(outdir),str(docx_path)]
        subprocess.run(cmd,check=True,stdout=subprocess.PIPE,stderr=subprocess.PIPE,env=env)
        made=outdir/(docx_path.stem+'.pdf')
        if not made.exists(): raise RuntimeError('LibreOffice did not create PDF')
        shutil.copy2(made,pdf_path)


def main():
    ap=argparse.ArgumentParser()
    ap.add_argument('report_json')
    ap.add_argument('-o','--output',required=True,help='Output PDF path')
    ap.add_argument('--keep-docx',action='store_true')
    ap.add_argument('--docx-output')
    args=ap.parse_args()
    data=json.loads(Path(args.report_json).read_text(encoding='utf-8'))
    pdf=Path(args.output); pdf.parent.mkdir(parents=True,exist_ok=True)
    if args.keep_docx:
        docx=Path(args.docx_output) if args.docx_output else pdf.with_suffix('.docx')
        build_docx(data,docx); convert_to_pdf(docx,pdf)
    else:
        with tempfile.TemporaryDirectory(prefix='cte_report_') as td:
            docx=Path(td)/'report.docx'; build_docx(data,docx); convert_to_pdf(docx,pdf)
    print(pdf)

if __name__=='__main__': main()
