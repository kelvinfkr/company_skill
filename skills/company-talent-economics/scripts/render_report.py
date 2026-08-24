#!/usr/bin/env python3
"""Typeset a Company Talent Economics report from structured JSON, in the report's own language.

Pipeline: report.json -> DOCX (python-docx) -> PDF (LibreOffice). Markdown is never the layout source.

Language comes from report.json meta.language (or --language). Section titles, table headers,
claim-type labels, number grouping and the page footer come from locales/<tag>.json; fonts are
chosen for the locale's script from what is actually installed; Arabic and Hebrew render RTL.
Analytical prose is whatever the agent wrote into the JSON and is passed through untouched.
"""
from __future__ import annotations
import argparse, json, os, shutil, subprocess, sys, tempfile
from pathlib import Path
from typing import Any, Iterable

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt, RGBColor

sys.path.insert(0, str(Path(__file__).resolve().parent))
import i18n

INK = "18212B"
MUTED = "667085"
LIGHT = "F2F4F7"
LINE = "D0D5DD"
DARK = "101828"
ACCENT = "344054"
SOFT = "F8F9FB"

# Set by build_docx() once the report's locale is known.
LOC: dict = {}
FONT_MAIN = "Noto Sans"
FONT_SERIF = "Noto Serif"
RTL = False


def T(path, default=""):
    return i18n.t(LOC, path, default)


def section(key):
    """(title, note) for a report section in the report's language."""
    v = T(f"sections.{key}", [key, ""])
    return (v + ["", ""])[:2] if isinstance(v, list) else (str(v), "")


def localize_claim(value):
    """Canonical claim-type keys written by the pipeline become display labels here;
    text the agent already wrote in its own language passes through."""
    if isinstance(value, str) and value in LOC.get("claim_types", {}):
        return LOC["claim_types"][value]
    return value


# ----------------------------------------------------------------- docx helpers

def set_cell_shading(cell, fill: str):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        tcPr.append(shd)
    shd.set(qn('w:fill'), fill)


def set_cell_border(cell, color=LINE, size=4):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = tcPr.first_child_found_in('w:tcBorders')
    if borders is None:
        borders = OxmlElement('w:tcBorders')
        tcPr.append(borders)
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        tag = 'w:' + edge
        el = borders.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            borders.append(el)
        el.set(qn('w:val'), 'single')
        el.set(qn('w:sz'), str(size))
        el.set(qn('w:color'), color)


def set_repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr()
    h = OxmlElement('w:tblHeader')
    h.set(qn('w:val'), 'true')
    trPr.append(h)
    set_row_cant_split(row)


def set_row_cant_split(row):
    """Move a whole row to the next page rather than splitting it: a row broken mid-cell leaves an
    orphan fragment that reads as an empty table."""
    trPr = row._tr.get_or_add_trPr()
    if trPr.find(qn('w:cantSplit')) is None:
        trPr.append(OxmlElement('w:cantSplit'))


def set_cell_margins(cell, top=80, start=90, bottom=80, end=90):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in('w:tcMar')
    if tcMar is None:
        tcMar = OxmlElement('w:tcMar')
        tcPr.append(tcMar)
    for m, v in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = tcMar.find(qn('w:' + m))
        if node is None:
            node = OxmlElement('w:' + m)
            tcMar.append(node)
        node.set(qn('w:w'), str(v))
        node.set(qn('w:type'), 'dxa')


def apply_bidi(paragraph):
    """Mark a paragraph right-to-left so Word/LibreOffice lay it out correctly."""
    if not RTL:
        return
    pPr = paragraph._p.get_or_add_pPr()
    if pPr.find(qn('w:bidi')) is None:
        pPr.append(OxmlElement('w:bidi'))
    if paragraph.alignment is None:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT


def apply_table_bidi(table):
    """Mirror column order for RTL so the first header cell sits on the right."""
    if not RTL:
        return
    tblPr = table._tbl.tblPr
    if tblPr.find(qn('w:bidiVisual')) is None:
        tblPr.append(OxmlElement('w:bidiVisual'))


def text_run(p, text, *, bold=False, size=None, color=None, font=None):
    font = font or FONT_MAIN
    r = p.add_run("" if text is None else str(text))
    r.bold = bold
    r.font.name = font
    r._element.rPr.rFonts.set(qn('w:eastAsia'), font)
    r._element.rPr.rFonts.set(qn('w:cs'), font)
    if size:
        r.font.size = Pt(size)
    if color:
        r.font.color.rgb = RGBColor.from_string(color)
    if RTL:
        rtl = OxmlElement('w:rtl')
        rtl.set(qn('w:val'), 'true')
        r._element.rPr.append(rtl)
    apply_bidi(p)
    return r


def header_footer_font(paragraph, size=8, color=MUTED):
    for run in paragraph.runs:
        run.font.name = FONT_MAIN
        run._element.rPr.rFonts.set(qn('w:eastAsia'), FONT_MAIN)
        run.font.size = Pt(size)
        run.font.color.rgb = RGBColor.from_string(color)


def add_page_number(paragraph):
    """Render the locale's page template, placing the PAGE field where {n} sits."""
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT if RTL else WD_ALIGN_PARAGRAPH.RIGHT
    template = T("cover.page", "Page {n}")
    before, _, after = template.partition("{n}")
    if before:
        paragraph.add_run(before)
    run = paragraph.add_run()
    begin = OxmlElement('w:fldChar'); begin.set(qn('w:fldCharType'), 'begin')
    instr = OxmlElement('w:instrText'); instr.set(qn('xml:space'), 'preserve'); instr.text = 'PAGE'
    end = OxmlElement('w:fldChar'); end.set(qn('w:fldCharType'), 'end')
    run._r.append(begin); run._r.append(instr); run._r.append(end)
    if after:
        paragraph.add_run(after)
    header_footer_font(paragraph)
    apply_bidi(paragraph)


def configure_doc(doc: Document):
    sec = doc.sections[0]
    sec.page_width = Mm(210)
    sec.page_height = Mm(297)
    sec.different_first_page_header_footer = True
    sec.top_margin = Mm(18); sec.bottom_margin = Mm(17)
    sec.left_margin = Mm(18); sec.right_margin = Mm(18)
    sec.header_distance = Mm(8); sec.footer_distance = Mm(8)
    if RTL:
        sectPr = sec._sectPr
        if sectPr.find(qn('w:bidi')) is None:
            sectPr.append(OxmlElement('w:bidi'))

    styles = doc.styles
    normal = styles['Normal']
    normal.font.name = FONT_MAIN
    normal._element.rPr.rFonts.set(qn('w:eastAsia'), FONT_MAIN)
    normal._element.rPr.rFonts.set(qn('w:cs'), FONT_MAIN)
    normal.font.size = Pt(9.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(4.5)
    normal.paragraph_format.line_spacing = 1.22

    for nm, size, before, after, color in [
            ('Title', 25, 0, 8, DARK), ('Subtitle', 10.5, 0, 8, MUTED),
            ('Heading 1', 16, 16, 8, DARK), ('Heading 2', 12, 10, 5, ACCENT),
            ('Heading 3', 10.5, 8, 4, ACCENT)]:
        st = styles[nm]
        st.font.name = FONT_MAIN
        st._element.rPr.rFonts.set(qn('w:eastAsia'), FONT_MAIN)
        st._element.rPr.rFonts.set(qn('w:cs'), FONT_MAIN)
        st.font.size = Pt(size)
        st.font.color.rgb = RGBColor.from_string(color)
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.keep_with_next = True


# ------------------------------------------------------------------- components

def add_cover(doc: Document, data: dict):
    meta = data.get('meta', {})
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(78)
    text_run(p, meta.get('company', 'Company'), bold=True, size=27, color=DARK, font=FONT_SERIF)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    text_run(p, T("cover.report_title"), bold=True, size=19, color=DARK)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    text_run(p, meta.get('subtitle') or T("cover.subtitle_default"), size=10.5, color=MUTED)

    doc.add_paragraph('')
    dash = T("labels.dash", "-")
    rows = [(T("cover.reference_date"), meta.get('reference_date', dash)),
            (T("cover.entity_type"), meta.get('entity_type', dash)),
            (T("cover.scope"), meta.get('scope', dash)),
            (T("cover.threshold"), meta.get('threshold', dash)),
            (T("cover.confidence"), meta.get('confidence', dash))]
    t = doc.add_table(rows=0, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    apply_table_bidi(t)
    for k, v in rows:
        c = t.add_row().cells
        c[0].width = Cm(3.6); c[1].width = Cm(9.2)
        set_cell_shading(c[0], LIGHT)
        for idx, txt in enumerate((k, v)):
            c[idx].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(c[idx], 120, 140, 120, 140)
            p = c[idx].paragraphs[0]; p.paragraph_format.space_after = Pt(0)
            text_run(p, txt, bold=(idx == 0), size=9.2, color=INK if idx else ACCENT)
    doc.add_paragraph('')
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    text_run(p, T("cover.method_note"), size=8.2, color=MUTED)
    if data.get('meta', {}).get('provisional_note'):
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        text_run(p, data['meta']['provisional_note'], size=8.2, color=DARK, )
    doc.add_page_break()


def add_header_footer(doc: Document, company: str):
    sec = doc.sections[0]
    hdr = sec.header.paragraphs[0]
    hdr.alignment = WD_ALIGN_PARAGRAPH.RIGHT if RTL else WD_ALIGN_PARAGRAPH.LEFT
    text_run(hdr, f'{company}  |  {T("cover.header_suffix")}', size=7.8, color=MUTED)
    add_page_number(sec.footer.paragraphs[0])


def add_section_title(doc, num, title, note=None):
    p = doc.add_paragraph(style='Heading 1')
    text_run(p, f'{num:02d}', bold=True, size=9.5, color=MUTED)
    text_run(p, f'  {title}', bold=True, size=16, color=DARK)
    if note:
        q = doc.add_paragraph(); q.paragraph_format.space_after = Pt(7)
        text_run(q, note, size=8.8, color=MUTED)


def add_callout(doc, text):
    t = doc.add_table(rows=1, cols=1)
    t.autofit = False; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    apply_table_bidi(t)
    cell = t.cell(0, 0)
    set_cell_shading(cell, SOFT); set_cell_margins(cell, 160, 180, 160, 180)
    p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(0)
    text_run(p, text, bold=True, size=10, color=DARK)
    doc.add_paragraph('').paragraph_format.space_after = Pt(2)


def add_kpi_cards(doc, metrics: list):
    if not metrics:
        return
    for start in range(0, len(metrics), 3):
        batch = metrics[start:start + 3]
        t = doc.add_table(rows=1, cols=len(batch))
        t.autofit = False; t.alignment = WD_TABLE_ALIGNMENT.CENTER
        apply_table_bidi(t)
        for i, m in enumerate(batch):
            c = t.cell(0, i); c.width = Cm(5.2)
            set_cell_shading(c, SOFT); set_cell_margins(c, 160, 150, 140, 150)
            p = c.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(3)
            text_run(p, m.get('value', T("labels.dash", "-")), bold=True, size=15, color=DARK)
            p = c.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            text_run(p, m.get('label', ''), size=8.2, color=MUTED)
        doc.add_paragraph('').paragraph_format.space_after = Pt(1)


def add_table(doc, headers: list, rows: Iterable[Iterable[Any]], widths: list | None = None, font_size=8.0):
    rows = list(rows)
    if not headers:
        return
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    apply_table_bidi(t)
    if widths is None:
        widths = [16.0 / len(headers)] * len(headers)
    hdr = t.rows[0]
    set_repeat_table_header(hdr)
    for i, h in enumerate(headers):
        c = hdr.cells[i]; c.width = Cm(widths[i])
        set_cell_shading(c, ACCENT); set_cell_margins(c, 100, 100, 100, 100)
        p = c.paragraphs[0]; p.paragraph_format.space_after = Pt(0)
        text_run(p, h, bold=True, size=font_size, color='FFFFFF')
    for ri, row in enumerate(rows):
        new_row = t.add_row()
        set_row_cant_split(new_row)
        cells = new_row.cells
        for i, val in enumerate(row):
            if i >= len(cells):
                break
            c = cells[i]; c.width = Cm(widths[i])
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(c, 90, 100, 90, 100); set_cell_border(c)
            if ri % 2 == 1:
                set_cell_shading(c, 'FAFAFA')
            p = c.paragraphs[0]; p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            text_run(p, '' if val is None else localize_claim(val), size=font_size, color=INK)
    doc.add_paragraph('').paragraph_format.space_after = Pt(1)


def rows_by_keys(items, keys):
    return [[x.get(k, '') for k in keys] for x in items or []]


def add_bullets(doc, items):
    for item in items or []:
        p = doc.add_paragraph(style='Normal')
        p.paragraph_format.left_indent = Cm(0.45)
        p.paragraph_format.first_line_indent = Cm(-0.28)
        text_run(p, '• ', bold=True, color=ACCENT)
        text_run(p, item, size=9.2)


def add_subhead(doc, text):
    p = doc.add_paragraph(style='Heading 2')
    text_run(p, text, bold=True, size=12, color=ACCENT)


# ----------------------------------------------------------------------- report

def build_docx(data: dict, out_docx: Path, language: str | None = None):
    global LOC, FONT_MAIN, FONT_SERIF, RTL
    lang = language or data.get('meta', {}).get('language')
    if not lang:
        # Older bundles carry no meta.language: fall back to the company name's script rather
        # than dressing a Chinese report in English chrome.
        guess = i18n.detect_language(data.get('meta', {}).get('company', ''))
        lang = guess["language"]
        print(f"[note] report.json has no meta.language; using '{lang}' ({guess['signal']}). "
              f"Set meta.language to be explicit.", file=sys.stderr)
    LOC = i18n.load_locale(lang)
    RTL = i18n.is_rtl(LOC)
    FONT_MAIN, FONT_SERIF = i18n.fonts(LOC)
    private_font_dir = os.environ.get('CTE_FONT_DIR')
    private_fonts = Path(private_font_dir) if private_font_dir else None
    if (private_fonts and private_fonts.is_dir()
            and LOC.get('meta', {}).get('script') in ('hans', 'hant', 'jpan', 'kore')
            and any(private_fonts.glob('NotoSansCJK*.otf'))):
        FONT_MAIN = FONT_SERIF = 'Noto Sans CJK SC'
        warn = None
    else:
        warn = i18n.font_warning(LOC)
    if warn:
        print(f"[!!] {warn}", file=sys.stderr)
    if LOC.get("_chrome_is_fallback"):
        print(f"[note] no locale pack for '{LOC['_requested']}'; section titles and table headers "
              f"render in {LOC['_chrome']}. Prose stays as written. "
              f"Add locales/{LOC['_requested']}.json to fix that.", file=sys.stderr)

    doc = Document()
    configure_doc(doc)
    company = data.get('meta', {}).get('company', 'Company')
    add_cover(doc, data)
    add_header_footer(doc, company)

    title, note = section("executive")
    add_section_title(doc, 1, title, note)
    ex = data.get('executive', {})
    if ex.get('diagnosis'):
        add_callout(doc, ex['diagnosis'])
    add_kpi_cards(doc, ex.get('metrics', []))

    title, note = section("charter")
    add_section_title(doc, 2, title, note)
    charter = data.get('charter') or data.get('constitution', {})
    add_table(doc, T("tables.charter_summary"),
              rows_by_keys(charter.get('summary', []), ['item', 'current', 'historical', 'evidence']),
              [3.5, 4.5, 5.0, 3.0], 7.8)
    if charter.get('timeline'):
        add_subhead(doc, T("labels.charter_timeline"))
        add_table(doc, T("tables.charter_timeline"),
                  rows_by_keys(charter['timeline'], ['date', 'event', 'economic', 'control', 'status']),
                  [2.2, 4.2, 3.4, 4.1, 2.1], 7.3)

    doc.add_page_break()
    title, note = section("business")
    add_section_title(doc, 3, title, note)
    bus = data.get('business', {})
    add_table(doc, T("tables.business_metrics"),
              rows_by_keys(bus.get('metrics', []), ['metric', 'value', 'type', 'evidence']),
              [4.4, 3.2, 3.0, 5.4], 8.0)
    if bus.get('profit_pools'):
        add_subhead(doc, T("labels.profit_pools"))
        add_table(doc, T("tables.profit_pools"),
                  rows_by_keys(bus['profit_pools'], ['segment', 'revenue', 'gross_profit', 'growth', 'dependency']),
                  [3.2, 2.7, 3.1, 2.0, 5.0], 7.7)

    title, note = section("talent_pnl")
    add_section_title(doc, 4, title, note)
    add_table(doc, T("tables.talent_pnl"),
              rows_by_keys(data.get('talent_pnl', []), ['profit_pool', 'process', 'role', 'dependency', 'why']),
              [3.2, 3.1, 3.4, 2.1, 4.2], 7.5)

    doc.add_page_break()
    title, note = section("organization")
    add_section_title(doc, 5, title, note)
    add_table(doc, T("tables.organization"),
              rows_by_keys(data.get('organization', []),
                           ['native', 'track', 'scope', 'function', 'state', 'confidence']),
              [3.6, 2.1, 2.5, 3.4, 2.5, 1.9], 7.3)

    title, note = section("compensation")
    add_section_title(doc, 6, title, note)
    add_table(doc, T("tables.compensation"),
              rows_by_keys(data.get('compensation', []),
                           ['bucket', 'headcount', 'cash', 'equity', 'total', 'probability', 'confidence']),
              [3.1, 1.8, 2.5, 2.5, 2.7, 2.1, 1.8], 7.1)

    doc.add_page_break()
    title, note = section("high_compensation")
    add_section_title(doc, 7, title, note)
    high = data.get('high_compensation', {})
    if high.get('headline'):
        add_callout(doc, T("labels.high_comp_headline", "{v}").format(v=high['headline']))
    add_table(doc, T("tables.high_compensation"),
              rows_by_keys(high.get('rows', []), ['bucket', 'count', 'contribution', 'uncertainty']),
              [4.0, 3.0, 2.4, 6.0], 7.8)

    title, note = section("role_pay")
    add_section_title(doc, 8, title, note)
    add_table(doc, T("tables.role_pay"),
              rows_by_keys(data.get('role_pay', []),
                           ['role', 'market', 'economic', 'internal', 'defensible', 'diagnostic']),
              [3.4, 2.5, 2.8, 2.4, 2.6, 2.3], 7.2)

    doc.add_page_break()
    title, note = section("career")
    add_section_title(doc, 9, title, note)
    car = data.get('career', {})
    if car.get('internal_path'):
        add_subhead(doc, T("labels.internal_path"))
        add_bullets(doc, car['internal_path'])
    if car.get('external_path'):
        add_subhead(doc, T("labels.external_path"))
        add_bullets(doc, car['external_path'])
    add_table(doc, T("tables.career_qa"),
              [[T("labels.first_threshold_state"), car.get('first_threshold_state', '')],
               [T("labels.promotion_tenure"), car.get('promotion_tenure', '')],
               [T("labels.senior_hire_mix"), car.get('senior_hire_mix', '')]],
              [5.0, 11.0], 8.0)

    title, note = section("governance")
    add_section_title(doc, 10, title, note)
    add_table(doc, T("tables.governance"),
              rows_by_keys(data.get('governance', []), ['metric', 'value', 'evidence']),
              [5.0, 6.0, 5.0], 7.9)

    title, note = section("sensitivity")
    add_section_title(doc, 11, title, note)
    add_bullets(doc, data.get('sensitivity', []))

    title, note = section("evidence")
    add_section_title(doc, 12, title, note)
    add_table(doc, T("tables.evidence"),
              rows_by_keys(data.get('evidence', []), ['id', 'source', 'date', 'status', 'supports']),
              [1.4, 5.1, 2.2, 2.0, 5.3], 7.1)

    doc.save(out_docx)


def convert_to_pdf(docx_path: Path, pdf_path: Path):
    with tempfile.TemporaryDirectory(prefix='cte_lo_') as td:
        env = os.environ.copy()
        env['HOME'] = td
        private_font_dir = env.get('CTE_FONT_DIR')
        if private_font_dir and Path(private_font_dir).is_dir():
            private_fonts = Path(td) / '.fonts'
            shutil.copytree(private_font_dir, private_fonts)
            env['SAL_PRIVATE_FONTPATH'] = str(private_fonts)
            fc_cache = shutil.which('fc-cache')
            if fc_cache:
                subprocess.run([fc_cache, '-f', str(private_fonts)], check=True,
                               stdout=subprocess.PIPE, stderr=subprocess.PIPE, env=env)
        outdir = Path(td) / 'out'
        outdir.mkdir()
        exe = next((e for e in ('libreoffice', 'soffice',
                                '/Applications/LibreOffice.app/Contents/MacOS/soffice',
                                r'C:\Program Files\LibreOffice\program\soffice.exe')
                    if shutil.which(e) or os.path.exists(e)), None)
        if exe is None:
            raise SystemExit('LibreOffice not found. Install it (apt: libreoffice-writer; '
                             'brew: --cask libreoffice) or run with --docx-only and convert elsewhere.')
        subprocess.run([exe, '--headless', '--convert-to', 'pdf', '--outdir', str(outdir), str(docx_path)],
                       check=True, stdout=subprocess.PIPE, stderr=subprocess.PIPE, env=env)
        made = outdir / (docx_path.stem + '.pdf')
        if not made.exists():
            raise RuntimeError('LibreOffice did not create PDF')
        shutil.copy2(made, pdf_path)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('report_json')
    ap.add_argument('-o', '--output', required=True, help='Output PDF path')
    ap.add_argument('--language', help='override meta.language')
    ap.add_argument('--keep-docx', action='store_true')
    ap.add_argument('--docx-output')
    ap.add_argument('--docx-only', action='store_true',
                    help='Write DOCX only; skip PDF conversion (no LibreOffice needed)')
    args = ap.parse_args()
    data = json.loads(Path(args.report_json).read_text(encoding='utf-8'))
    pdf = Path(args.output)
    pdf.parent.mkdir(parents=True, exist_ok=True)
    if args.docx_only:
        docx = Path(args.docx_output) if args.docx_output else pdf.with_suffix('.docx')
        build_docx(data, docx, args.language)
        print(docx)
        return
    if args.keep_docx:
        docx = Path(args.docx_output) if args.docx_output else pdf.with_suffix('.docx')
        build_docx(data, docx, args.language)
        convert_to_pdf(docx, pdf)
    else:
        with tempfile.TemporaryDirectory(prefix='cte_report_') as td:
            docx = Path(td) / 'report.docx'
            build_docx(data, docx, args.language)
            convert_to_pdf(docx, pdf)
    print(pdf)


if __name__ == '__main__':
    main()
